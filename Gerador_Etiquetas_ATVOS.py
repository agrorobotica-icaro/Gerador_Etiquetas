################################### ICARO ALVES RAMIRO ########################################################################
# Aplicação para gerar etiquetas e acompanhar o planejamento semanal de análises.
# Execute o programa com: streamlit run nome_do_arquivo.py
# Instale previamente as dependências: streamlit, openpyxl e reportlab.
# Na aba "Gerador de Etiquetas", selecione as planilhas ETIQUETAS e QUIMICO_SIMPLES.
# Cada célula preenchida na planilha ETIQUETAS gera uma etiqueta.
# Na aba "Planejamento semanal", envie uma ou mais planilhas padronizadas no formato XLSX.
# As planilhas de planejamento devem conter Ordem_de_Serviço e as colunas de análises previstas.
# O resumo considera apenas linhas com Ordem_de_Serviço no padrão OS_número_talhão_remessa_tipo.
# As linhas são agrupadas por número da OS, remessa e tipo; o talhão não aparece no resumo.
# Cada análise mostra a quantidade de marcações X; a coluna Total mostra o total de linhas.

import io
import re
import zipfile
from collections import defaultdict
from pathlib import Path

import streamlit as st
from openpyxl import load_workbook
from reportlab.lib.units import cm, mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas


st.set_page_config(
    page_title="Gerador de Etiquetas ATVOS",
    layout="centered",
)

st.markdown(
    """
    <style>
    .main { background-color: #f5f7fa; }
    .block-container {
        padding-top: 2rem;
        padding-bottom: 2rem;
        padding-left: 3rem;
        padding-right: 3rem;
    }
    .titulo-app {
        font-size: 38px;
        font-weight: 700;
        color: #0F172A;
        margin-bottom: 0px;
    }
    .subtitulo-app {
        font-size: 18px;
        color: #475569;
        margin-top: 0px;
        margin-bottom: 0px;
    }
    .stButton > button, .stDownloadButton > button {
        width: 100%;
        color: white;
        border-radius: 12px;
        border: none;
        height: 50px;
        font-size: 16px;
        font-weight: 600;
    }
    .stButton > button { background-color: #166534; }
    .stButton > button:hover {
        background-color: #14532d;
        color: white;
    }
    .stDownloadButton > button { background-color: #0F172A; }
    .stDownloadButton > button:hover {
        background-color: #1E293B;
        color: white;
    }
    .stAlert { border-radius: 12px; }
    hr { margin-top: 10px; margin-bottom: 25px; }
    </style>
    """,
    unsafe_allow_html=True,
)

st.markdown(
    """
    <div class="titulo-app">Gerador de Etiquetas ATVOS</div>
    <div class="subtitulo-app">Laboratório Químico • Agrorobótica</div>
    """,
    unsafe_allow_html=True,
)
st.divider()


# Configuração das etiquetas
LARGURA_PAGINA = 10 * cm
ALTURA_PAGINA = 1.5 * cm
ETIQUETAS_POR_PAGINA = 4
LARGURA_ETIQUETA = LARGURA_PAGINA / ETIQUETAS_POR_PAGINA
ESPACAMENTO_LINHAS = 1.0 * mm
TAMANHO_FONTE = 10

PARAMETROS = ("k", "m", "u", "s", "p", "c", "t", "r", "CHN")
PADRAO_ETIQUETA = re.compile(
    r"^\s*([A-Za-z]+)\s+(\d+)\|(\d+)\|([^|]+)\|([^|]+)\s*$"
)
PADRAO_OS = re.compile(
    r"^\s*OS[_\s-](\d+)_([^_]+)_([^_]+)_([^_]+)\s*$",
    re.IGNORECASE,
)
PADRAO_OS_PLANEJAMENTO = re.compile(
    r"^\s*OS[_\s-](\d+)_([^_]+)_([^_]+)_([^_]+)\s*$",
    re.IGNORECASE,
)

ANALISES_PLANEJAMENTO = (
    "K_P_Mehlich",
    "Macro",
    "Micro",
    "MO",
    "pH_CaCl2",
    "pH_H2O",
    "P_Resina",
    "S_ICP",
    "S_Turbidimetria",
    "Textura",
)


def normalizar_cabecalho(valor):
    return str(valor).strip() if valor is not None else ""


def ler_etiquetas(arquivo):
    """Lê as etiquetas preenchidas, preservando a ordem das colunas e linhas."""
    arquivo.seek(0)
    workbook = load_workbook(arquivo, read_only=True, data_only=True)
    sheet = workbook.active
    linhas = sheet.iter_rows(values_only=True)
    cabecalhos = [normalizar_cabecalho(v) for v in next(linhas, ())]
    indices = {nome: i for i, nome in enumerate(cabecalhos) if nome}
    ausentes = [sigla for sigla in PARAMETROS if sigla not in indices]
    if ausentes:
        workbook.close()
        raise ValueError(
            "A planilha ETIQUETAS não contém as colunas: " + ", ".join(ausentes)
        )

    etiquetas = []
    for numero_linha, linha in enumerate(linhas, start=2):
        for sigla in PARAMETROS:
            valor = linha[indices[sigla]] if indices[sigla] < len(linha) else None
            if valor is None or str(valor).strip() == "":
                continue
            correspondencia = PADRAO_ETIQUETA.match(str(valor))
            if not correspondencia:
                workbook.close()
                raise ValueError(
                    f"Formato inválido na planilha ETIQUETAS, linha {numero_linha}, "
                    f"coluna {sigla}: {valor!r}. Esperado: "
                    f"'{sigla} 003|331|R1|F'."
                )

            sigla_lida, sid, os_numero, remessa, tipo = correspondencia.groups()
            if sigla_lida.lower() != sigla.lower():
                workbook.close()
                raise ValueError(
                    f"A etiqueta {valor!r} está na coluna {sigla}, mas começa com "
                    f"'{sigla_lida}'."
                )

            etiquetas.append(
                {
                    "sigla": sigla,
                    "sid": int(sid),
                    "sid_texto": sid.zfill(3),
                    "os": os_numero,
                    "remessa": remessa.strip(),
                    "tipo": tipo.strip(),
                    "linha1": f"{sigla} {sid.zfill(3)}",
                    "linha3": f"{remessa.strip()}|{tipo.strip()}",
                }
            )
    workbook.close()
    return etiquetas


def ler_mapa_talhao(arquivo):
    """Mapeia (OS, remessa, tipo, SID) para Talhão_AG."""
    arquivo.seek(0)
    workbook = load_workbook(arquivo, read_only=True, data_only=True)
    sheet = workbook.active
    linhas = sheet.iter_rows(values_only=True)
    cabecalhos = [normalizar_cabecalho(v) for v in next(linhas, ())]
    indices = {nome: i for i, nome in enumerate(cabecalhos) if nome}
    obrigatorias = ("Ordem_de_Serviço", "Talhão_AG", "SID")
    ausentes = [nome for nome in obrigatorias if nome not in indices]
    if ausentes:
        workbook.close()
        raise ValueError(
            "A planilha QUIMICO_SIMPLES não contém as colunas: "
            + ", ".join(ausentes)
        )

    mapa = {}
    for linha in linhas:
        ordem = linha[indices["Ordem_de_Serviço"]]
        talhao_ag = linha[indices["Talhão_AG"]]
        sid = linha[indices["SID"]]
        if ordem in (None, "") or talhao_ag in (None, "") or sid in (None, ""):
            continue
        os_match = PADRAO_OS.match(str(ordem))
        if not os_match:
            continue
        os_numero, _talhao_do_nome, remessa, tipo = os_match.groups()
        chave = (os_numero, remessa.upper(), tipo.upper(), int(sid))
        valor = str(talhao_ag).strip()
        if chave in mapa and mapa[chave] != valor:
            workbook.close()
            raise ValueError(
                "A planilha QUIMICO_SIMPLES tem mais de um Talhão_AG para "
                f"OS {os_numero}, {remessa}, {tipo}, SID {int(sid)} "
                f"(valores {mapa[chave]!r} e {valor!r})."
            )
        mapa[chave] = valor
    workbook.close()
    return mapa


def associar_talhoes(etiquetas, mapa_talhao):
    sem_correspondencia = []
    for etiqueta in etiquetas:
        chave = (
            etiqueta["os"],
            etiqueta["remessa"].upper(),
            etiqueta["tipo"].upper(),
            etiqueta["sid"],
        )
        talhao_ag = mapa_talhao.get(chave)
        if talhao_ag is None:
            sem_correspondencia.append(
                f'{etiqueta["sigla"]} {etiqueta["sid_texto"]} | '
                f'OS {etiqueta["os"]} | {etiqueta["remessa"]}|{etiqueta["tipo"]}'
            )
        else:
            etiqueta["linha2"] = f'OS{etiqueta["os"]}|{talhao_ag}'

    if sem_correspondencia:
        exemplos = "\n".join(f"• {item}" for item in sem_correspondencia[:10])
        complemento = (
            f"\n... e mais {len(sem_correspondencia) - 10} etiqueta(s)."
            if len(sem_correspondencia) > 10
            else ""
        )
        raise ValueError(
            f"Não encontrei Talhão_AG para {len(sem_correspondencia)} etiqueta(s):\n"
            f"{exemplos}{complemento}"
        )
    return etiquetas


def registrar_fontes():
    """Usa Times New Roman no Windows e Helvetica como alternativa."""
    fonte_normal = "Helvetica"
    fonte_negrito = "Helvetica-Bold"
    fonte_normal_path = Path(r"C:\Windows\Fonts\times.ttf")
    fonte_negrito_path = Path(r"C:\Windows\Fonts\timesbd.ttf")
    if fonte_normal_path.exists() and fonte_negrito_path.exists():
        try:
            if "TimesNewRoman" not in pdfmetrics.getRegisteredFontNames():
                pdfmetrics.registerFont(TTFont("TimesNewRoman", str(fonte_normal_path)))
            if "TimesNewRomanBold" not in pdfmetrics.getRegisteredFontNames():
                pdfmetrics.registerFont(
                    TTFont("TimesNewRomanBold", str(fonte_negrito_path))
                )
            fonte_normal = "TimesNewRoman"
            fonte_negrito = "TimesNewRomanBold"
        except Exception:
            pass
    return fonte_normal, fonte_negrito


def desenhar_pagina(pdf, etiquetas_pagina, fontes):
    _, fonte_negrito = fontes
    passo_linha = TAMANHO_FONTE + ESPACAMENTO_LINHAS
    for indice, etiqueta in enumerate(etiquetas_pagina):
        x_centro = indice * LARGURA_ETIQUETA + LARGURA_ETIQUETA / 2
        y_centro = ALTURA_PAGINA / 2
        pdf.setFont(fonte_negrito, TAMANHO_FONTE)
        pdf.drawCentredString(x_centro, y_centro + passo_linha, etiqueta["linha1"])
        pdf.drawCentredString(x_centro, y_centro, etiqueta["linha2"])
        pdf.drawCentredString(x_centro, y_centro - passo_linha, etiqueta["linha3"])
    pdf.showPage()


def gerar_pdf(etiquetas, fontes):
    """Gera um PDF por parâmetro e tipo, com quatro etiquetas por página."""
    grupos = defaultdict(list)
    for etiqueta in etiquetas:
        grupos[(etiqueta["sigla"], etiqueta["tipo"])].append(etiqueta)
    arquivos_pdf = {}
    for (sigla, tipo), itens in grupos.items():
        buffer_pdf = io.BytesIO()
        pdf = canvas.Canvas(buffer_pdf, pagesize=(LARGURA_PAGINA, ALTURA_PAGINA))
        for inicio in range(0, len(itens), ETIQUETAS_POR_PAGINA):
            desenhar_pagina(pdf, itens[inicio : inicio + ETIQUETAS_POR_PAGINA], fontes)
        pdf.save()
        buffer_pdf.seek(0)
        arquivos_pdf[f"Etiquetas_{sigla}_{tipo}.pdf"] = buffer_pdf.getvalue()
    return arquivos_pdf


def gerar_zip(arquivos_pdf):
    buffer_zip = io.BytesIO()
    with zipfile.ZipFile(buffer_zip, "w", compression=zipfile.ZIP_DEFLATED) as zipf:
        for nome, conteudo in arquivos_pdf.items():
            zipf.writestr(nome, conteudo)
    buffer_zip.seek(0)
    return buffer_zip.getvalue()


def gerar_tabela_planejamento(arquivos):
    """Conta marcações X e total de linhas por OS/remessa/tipo."""
    totais = defaultdict(lambda: {"total": 0, **{a: 0 for a in ANALISES_PLANEJAMENTO}})
    problemas = []

    for arquivo in arquivos:
        arquivo.seek(0)
        workbook = load_workbook(arquivo, read_only=True, data_only=True)
        sheet = workbook.active
        linhas = sheet.iter_rows(values_only=True)
        cabecalhos = [normalizar_cabecalho(v) for v in next(linhas, ())]
        indices = {nome: i for i, nome in enumerate(cabecalhos) if nome}
        obrigatorias = ("Ordem_de_Serviço",) + ANALISES_PLANEJAMENTO
        ausentes = [nome for nome in obrigatorias if nome not in indices]
        if ausentes:
            problemas.append(f"{arquivo.name}: faltam as colunas {', '.join(ausentes)}")
            workbook.close()
            continue

        for numero_linha, linha in enumerate(linhas, start=2):
            ordem = linha[indices["Ordem_de_Serviço"]]
            if ordem is None or not str(ordem).strip():
                continue
            correspondencia = PADRAO_OS_PLANEJAMENTO.match(str(ordem))
            if not correspondencia:
                # Linhas sem identificação padronizada não entram no total.
                continue
            numero_os, _talhao, remessa, tipo = correspondencia.groups()
            chave_os = (numero_os, remessa.strip().upper(), tipo.strip().upper())
            grupo = totais[chave_os]
            grupo["total"] += 1
            for analise in ANALISES_PLANEJAMENTO:
                valor = linha[indices[analise]] if indices[analise] < len(linha) else None
                if str(valor).strip().upper() == "X":
                    grupo[analise] += 1
        workbook.close()

    if problemas:
        raise ValueError("\n".join(problemas))
    if not totais:
        return []

    tabela = []
    for chave_os in sorted(totais, key=lambda chave: (int(chave[0]), chave[1], chave[2])):
        numero_os, remessa, tipo = chave_os
        grupo = totais[chave_os]
        total = grupo["total"]
        linha = {
            "Ordem de Serviço": f"{numero_os}{remessa}{tipo}",
            "Total": total,
        }
        for analise in ANALISES_PLANEJAMENTO:
            linha[analise] = grupo[analise]
        tabela.append(linha)
    return tabela


aba_etiquetas, aba_planejamento = st.tabs(
    ["Gerador de Etiquetas", "Planejamento semanal"]
)

with aba_etiquetas:
    st.subheader("Selecione as planilhas")
    arquivo_etiquetas = st.file_uploader(
        "Planilha ETIQUETAS",
        type=["xlsx"],
        help="Planilha com os SIDs selecionados, organizados em colunas por parâmetro.",
        key="arquivo_etiquetas",
    )
    arquivo_quimico = st.file_uploader(
        "Planilha QUIMICO_SIMPLES",
        type=["xlsx"],
        help="Planilha usada para localizar o Talhão_AG de cada OS e SID.",
        key="arquivo_quimico",
    )
    st.caption(
        "Cada célula preenchida da planilha ETIQUETAS gera uma etiqueta. "
        "A identificação do Talhão_AG é feita pela OS, remessa, tipo de análise e SID."
    )

    if st.button("Gerar etiquetas", key="botao_gerar_etiquetas"):
        st.session_state.pop("etiquetas_zip", None)
        if arquivo_etiquetas is None or arquivo_quimico is None:
            st.error("Selecione as duas planilhas para continuar.")
        else:
            try:
                etiquetas = ler_etiquetas(arquivo_etiquetas)
                if not etiquetas:
                    st.error("Não encontrei etiquetas preenchidas na planilha ETIQUETAS.")
                else:
                    mapa_talhao = ler_mapa_talhao(arquivo_quimico)
                    etiquetas = associar_talhoes(etiquetas, mapa_talhao)
                    arquivos_pdf = gerar_pdf(etiquetas, registrar_fontes())
                    st.session_state["etiquetas_zip"] = gerar_zip(arquivos_pdf)
                    st.success(
                        f"Foram geradas {len(etiquetas)} etiquetas em "
                        f"{len(arquivos_pdf)} PDF(s)."
                    )
            except (ValueError, KeyError, IndexError) as erro:
                st.error(str(erro))
            except Exception as erro:
                st.error(f"Não foi possível gerar as etiquetas: {erro}")

    if "etiquetas_zip" in st.session_state:
        st.download_button(
            label="Baixar PDFs em ZIP",
            data=st.session_state["etiquetas_zip"],
            file_name="Etiquetas_ATVOS.zip",
            mime="application/zip",
            key="download_etiquetas_zip",
        )

with aba_planejamento:
    st.subheader("Planejamento semanal")
    planilhas_planejamento = st.file_uploader(
        "Selecione uma ou mais planilhas padronizadas",
        type=["xlsx"],
        accept_multiple_files=True,
        help=(
            "A tabela reúne as linhas de todas as planilhas e agrupa pela OS. "
            "A coluna Total mostra o número de linhas da OS/remessa/tipo; "
            "as colunas de análise mostram a quantidade de X."
        ),
        key="planilhas_planejamento",
    )
    st.caption(
        "Exemplo: 30 marcações em 100 linhas serão exibidas como 30 na análise e "
        "100 na coluna Total. "
        "Linhas sem Ordem_de_Serviço no padrão OS_331_A1_R1_F são ignoradas."
    )

    if planilhas_planejamento:
        try:
            tabela_planejamento = gerar_tabela_planejamento(planilhas_planejamento)
            if tabela_planejamento:
                st.dataframe(
                    tabela_planejamento,
                    hide_index=True,
                    use_container_width=True,
                )
            else:
                st.info("Não encontrei linhas com Ordem_de_Serviço padronizada.")
        except (ValueError, KeyError, IndexError) as erro:
            st.error(str(erro))
        except Exception as erro:
            st.error(f"Não foi possível montar o planejamento: {erro}")
    else:
        st.info("Envie as planilhas para montar o resumo por OS e análise.")
